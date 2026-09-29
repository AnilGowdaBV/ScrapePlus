import csv
import io
from collections.abc import Callable

import openpyxl
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from backend.app.database.base import Base
from backend.app.main import create_app
from scraper.people.models import PeopleSearchResult
from scraper.providers.base import PeopleSearchProvider

SEARCH_URL = "https://www.linkedin.com/search/results/people/?keywords=engineering"


class FakeProvider(PeopleSearchProvider):
    provider_name = "fake"

    def __init__(self, results: list[PeopleSearchResult]) -> None:
        self.results = results
        self.pages_processed = 1

    async def search(self, search_url: str, max_pages: int | None = None):
        for result in self.results:
            yield result


def client_for(factory: Callable[[], PeopleSearchProvider]) -> TestClient:
    engine = create_engine(
        "sqlite://",
        connect_args={"check_same_thread": False},
        poolclass=StaticPool,
    )
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine, expire_on_commit=False)
    return TestClient(
        create_app(provider_factory=factory, session_factory=session_factory)
    )


def test_export_csv_and_xlsx_endpoints() -> None:
    results = [
        PeopleSearchResult(
            person_name="Alice Smith",
            person_title="Software Architect",
            headline="Lead Architect at TechCorp",
            company_name="TechCorp",
            company_url="https://www.linkedin.com/company/techcorp",
            location="San Francisco, CA",
            linkedin_profile_url="https://www.linkedin.com/in/alice-smith/",
            raw_data={},
        ),
        PeopleSearchResult(
            person_name="Bob Jones",
            person_title="Data Scientist",
            headline="Data Scientist",
            company_name="DataCo",
            company_url=None,
            location="Austin, TX",
            linkedin_profile_url="https://www.linkedin.com/in/bob-jones/",
            raw_data={},
        ),
    ]
    with client_for(lambda: FakeProvider(results)) as client:
        created = client.post("/api/searches", json={"search_url": SEARCH_URL}).json()
        search_id = created["search_id"]

        # CSV Export
        csv_resp = client.get(f"/api/searches/{search_id}/export/csv")
        assert csv_resp.status_code == 200
        assert "text/csv" in csv_resp.headers["content-type"]
        assert f'filename="leads_search_{search_id}.csv"' in csv_resp.headers["content-disposition"]

        csv_text = csv_resp.content.decode("utf-8-sig")
        reader = list(csv.reader(io.StringIO(csv_text)))
        assert len(reader) == 3  # Header + 2 data rows
        assert reader[0] == ["Company", "Person Name", "Title / Headline", "Location", "LinkedIn Profile URL"]
        assert reader[1] == ["TechCorp", "Alice Smith", "Software Architect - Lead Architect at TechCorp", "San Francisco, CA", "https://www.linkedin.com/in/alice-smith"]
        assert reader[2] == ["DataCo", "Bob Jones", "Data Scientist", "Austin, TX", "https://www.linkedin.com/in/bob-jones"]

        # XLSX Export
        xlsx_resp = client.get(f"/api/searches/{search_id}/export/xlsx")
        assert xlsx_resp.status_code == 200
        assert "openxmlformats" in xlsx_resp.headers["content-type"]
        assert f'filename="leads_search_{search_id}.xlsx"' in xlsx_resp.headers["content-disposition"]

        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        assert len(rows) == 3
        assert list(rows[0]) == ["Company", "Person Name", "Title / Headline", "Location", "LinkedIn Profile URL"]
        assert list(rows[1]) == ["TechCorp", "Alice Smith", "Software Architect - Lead Architect at TechCorp", "San Francisco, CA", "https://www.linkedin.com/in/alice-smith"]
        assert list(rows[2]) == ["DataCo", "Bob Jones", "Data Scientist", "Austin, TX", "https://www.linkedin.com/in/bob-jones"]


def test_export_empty_results() -> None:
    with client_for(lambda: FakeProvider([])) as client:
        created = client.post("/api/searches", json={"search_url": SEARCH_URL}).json()
        search_id = created["search_id"]

        csv_resp = client.get(f"/api/searches/{search_id}/export/csv")
        assert csv_resp.status_code == 200
        reader = list(csv.reader(io.StringIO(csv_resp.content.decode("utf-8-sig"))))
        assert len(reader) == 1
        assert reader[0] == ["Company", "Person Name", "Title / Headline", "Location", "LinkedIn Profile URL"]

        xlsx_resp = client.get(f"/api/searches/{search_id}/export/xlsx")
        assert xlsx_resp.status_code == 200
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 1
        assert list(rows[0]) == ["Company", "Person Name", "Title / Headline", "Location", "LinkedIn Profile URL"]


def test_export_nonexistent_search_returns_404() -> None:
    with client_for(lambda: FakeProvider([])) as client:
        assert client.get("/api/searches/9999/export/csv").status_code == 404
        assert client.get("/api/searches/9999/export/xlsx").status_code == 404


def test_export_unicode_commas_quotes_and_newlines() -> None:
    results = [
        PeopleSearchResult(
            person_name='José "Pepe" Peña',
            person_title="VP, Engineering",
            headline="VP, Engineering\nCloud Systems",
            company_name="Acme, Inc.",
            company_url=None,
            location="München, Bayern, Germany",
            linkedin_profile_url="https://www.linkedin.com/in/jose-pepe/",
            raw_data={},
        ),
        PeopleSearchResult(
            person_name="山田 太郎",
            person_title="エンジニア",
            headline="シニアエンジニア",
            company_name="東京テクノロジー株式会社",
            company_url=None,
            location="東京, 日本",
            linkedin_profile_url="https://www.linkedin.com/in/yamada-taro/",
            raw_data={},
        ),
    ]
    with client_for(lambda: FakeProvider(results)) as client:
        created = client.post("/api/searches", json={"search_url": SEARCH_URL}).json()
        search_id = created["search_id"]

        # CSV checks
        csv_resp = client.get(f"/api/searches/{search_id}/export/csv")
        reader = list(csv.reader(io.StringIO(csv_resp.content.decode("utf-8-sig"))))
        assert len(reader) == 3
        # Commas and quotes preserved
        assert reader[1][0] == "Acme, Inc."
        assert reader[1][1] == 'José "Pepe" Peña'
        assert reader[1][3] == "München, Bayern, Germany"
        # Japanese Unicode preserved
        assert reader[2][0] == "東京テクノロジー株式会社"
        assert reader[2][1] == "山田 太郎"
        assert reader[2][3] == "東京, 日本"

        # XLSX checks
        xlsx_resp = client.get(f"/api/searches/{search_id}/export/xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 3
        assert rows[1][0] == "Acme, Inc."
        assert rows[1][1] == 'José "Pepe" Peña'
        assert rows[1][3] == "München, Bayern, Germany"
        assert rows[2][0] == "東京テクノロジー株式会社"
        assert rows[2][1] == "山田 太郎"
        assert rows[2][3] == "東京, 日本"


def test_export_null_fields_handled_cleanly() -> None:
    results = [
        PeopleSearchResult(
            person_name=None,
            person_title=None,
            headline=None,
            company_name=None,
            company_url=None,
            location=None,
            linkedin_profile_url=None,
            raw_data={},
        )
    ]
    with client_for(lambda: FakeProvider(results)) as client:
        created = client.post("/api/searches", json={"search_url": SEARCH_URL}).json()
        search_id = created["search_id"]

        csv_resp = client.get(f"/api/searches/{search_id}/export/csv")
        reader = list(csv.reader(io.StringIO(csv_resp.content.decode("utf-8-sig"))))
        assert len(reader) == 2
        assert reader[1] == ["", "", "", "", ""]

        xlsx_resp = client.get(f"/api/searches/{search_id}/export/xlsx")
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_resp.content))
        rows = list(wb.active.iter_rows(values_only=True))
        assert len(rows) == 2
        # None or "" handled safely
        assert all(cell is None or cell == "" for cell in rows[1])

